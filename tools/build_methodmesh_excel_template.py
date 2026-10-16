#!/usr/bin/env python3
"""Create the Excel-first MethodMesh NFC authoring template.

This is deliberately separate from the XLSForm compiler. Authors mark fields
with yes on the MethodMesh recipe sheet, copy the generated recipe and
canonical expression into the attestation rows, and then use the workbook as
an ordinary XLSForm.
"""

from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/Users/icrucrob/Downloads/ODK XLSForm Template.xlsx")
OUTPUT = ROOT / "app/src/main/java/com/example/methodmesh/modules/nfc/docs/example_odk_nfc_excel_authoring_template.xlsx"
SURVEY_START = 30
SURVEY_END = 468
RECIPE_START = 13


AUTH_FIELDS = [
    "methodmesh_execution_id", "methodmesh_method_id", "methodmesh_status",
    "credential_verified", "credential_verification_message", "credential_id",
    "credential_subject_id", "valid_until_iso", "pin_verified", "issuer_signature_valid",
    "issuer_trust_status", "issuer_key_id", "issuer_public_key_fingerprint_sha256",
    "tag_uid_hex", "verification_evidence_hash", "credential_envelope_hash",
    "credential_verified_time_iso", "methodmesh_full_json",
]

ATT_FIELDS = [
    "methodmesh_execution_id", "methodmesh_status", "attestation_id", "event_payload_hash",
    "attestation_schema_version", "commitment_recipe", "commitment_recipe_sha256",
    "verification_method", "verification_evidence_format", "verification_evidence_hash",
    "attestation_hash", "hash_algorithm", "public_key_id", "signature",
    "signature_algorithm", "trusted_timestamp_policy", "trusted_timestamp_status",
    "trusted_timestamp_authority", "trusted_timestamp_time_iso", "diagnostic_reason",
    "attestation_full_json",
]


def add_row(ws, row, type_, name, label="", *, hint="", required="", relevant="", appearance="", calculation="", constraint="", constraint_message="", readonly="", intent="", mm_commit=""):
    values = {
        "type": type_, "name": name, "label": label, "hint": hint, "required": required,
        "relevant": relevant, "appearance": appearance, "calculation": calculation,
        "constraint": constraint, "constraint_message": constraint_message,
        "readonly": readonly, "body::intent": intent, "mm_commit": mm_commit,
    }
    headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1) if ws.cell(1, col).value}
    for key, value in values.items():
        if key in headers:
            ws.cell(row, headers[key]).value = value


def ensure_survey_headers(ws):
    headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1) if ws.cell(1, col).value}
    for name in ("readonly", "body::intent"):
        if name not in headers:
            col = max(headers.values(), default=0) + 1
            ws.cell(1, col).value = name
            headers[name] = col
    return headers


def add_generic_scaffold(ws):
    headers = ensure_survey_headers(ws)
    # Keep the author's editable area stable so the recipe sheet can watch it.
    for r in range(2, 500):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None

    add_row(ws, 2, "calculate", "mm_workflow_instance_id", calculation="once(uuid())", appearance="hidden-answer")
    add_row(ws, 3, "calculate", "mm_study_id", calculation="''", appearance="hidden-answer")
    add_row(ws, 4, "calculate", "mm_form_id", calculation="''", appearance="hidden-answer")
    add_row(ws, 5, "calculate", "mm_form_version", calculation="''", appearance="hidden-answer")
    add_row(ws, 6, "calculate", "mm_timestamp_policy", calculation="'preferred'", appearance="hidden-answer")
    auth_intent = (
        "com.example.methodmesh.EXECUTE_METHOD(method_id='nfc_credential_verification',"
        "caller='odk',input_payload_mode='FULL',return_mode='flat',"
        "methodmesh_return_namespace='mm_auth')"
    )
    add_row(ws, 7, "begin group", "mm_authenticate_operator", "Authenticate operator",
            hint="Tap the NFC staff credential and enter its PIN.", appearance="field-list", intent=auth_intent)
    for offset, field in enumerate(AUTH_FIELDS, 8):
        add_row(ws, offset, "text", f"mm_auth_{field}", f"MethodMesh authentication return: {field}", appearance="hidden-answer", readonly="yes")
    add_row(ws, 8 + len(AUTH_FIELDS), "end group", "mm_authenticate_operator_end")
    add_row(ws, 9 + len(AUTH_FIELDS), "calculate", "mm_auth_ok",
            calculation="if(${mm_auth_methodmesh_status} = 'Succeeded' and ${mm_auth_credential_verified} = 'true' and ${mm_auth_pin_verified} = 'true' and ${mm_auth_issuer_signature_valid} = 'true', 'true', 'false')",
            appearance="hidden-answer")
    add_row(ws, 10 + len(AUTH_FIELDS), "note", "mm_auth_status_ok", "Credential verified",
            relevant="${mm_auth_ok} = 'true'")
    add_row(ws, 11 + len(AUTH_FIELDS), "note", "mm_auth_status_fail", "Authentication not accepted",
            hint="The study data remain locked until the NFC credential and PIN verify.",
            relevant="${mm_auth_methodmesh_status} != '' and ${mm_auth_ok} != 'true'")
    add_row(ws, SURVEY_START - 1, "begin group", "mm_study_data", "Study data",
            hint="Add ordinary study questions in rows 30–468. Mark inclusion on the MethodMesh recipe sheet.",
            relevant="${mm_auth_ok} = 'true'")
    add_row(ws, SURVEY_END + 1, "end group", "mm_study_data_end")

    r = SURVEY_END + 3
    add_row(ws, r, "calculate", "mm_credential_id_sha256", calculation="digest(string(${mm_auth_credential_id}), 'SHA-256', 'hex')", appearance="hidden-answer"); r += 1
    add_row(ws, r, "calculate", "mm_credential_subject_id_sha256", calculation="digest(string(${mm_auth_credential_subject_id}), 'SHA-256', 'hex')", appearance="hidden-answer"); r += 1
    add_row(ws, r, "calculate", "mm_canonical_commitment_live", calculation="PASTE_CANONICAL_EXPRESSION_FROM_METHODMESH_RECIPE_SHEET", appearance="hidden-answer"); r += 1
    add_row(ws, r, "calculate", "mm_current_event_payload_hash", calculation="digest(${mm_canonical_commitment_live}, 'SHA-256', 'hex')", appearance="hidden-answer"); r += 1
    add_row(ws, r, "select_one mm_yes", "mm_finalize_for_attestation", "Finalize these data for attestation?",
            required="yes", relevant="${mm_auth_ok} = 'true'", constraint=". = 'yes'",
            constraint_message="Select Yes to freeze the study data before attestation."); r += 1
    add_row(ws, r, "calculate", "mm_frozen_canonical_commitment", calculation="once(if(${mm_finalize_for_attestation} = 'yes', ${mm_canonical_commitment_live}, ''))", appearance="hidden-answer"); r += 1
    add_row(ws, r, "calculate", "mm_event_payload_hash", calculation="once(if(string-length(${mm_frozen_canonical_commitment}) > 0, digest(${mm_frozen_canonical_commitment}, 'SHA-256', 'hex'), ''))", appearance="hidden-answer"); r += 1
    add_row(ws, r, "calculate", "mm_commitment_recipe", calculation="PASTE_RECIPE_JSON_FROM_METHODMESH_RECIPE_SHEET", appearance="hidden-answer"); r += 1
    add_row(ws, r, "calculate", "mm_current_data_matches_frozen_commitment", calculation="if(${mm_current_event_payload_hash} = ${mm_event_payload_hash}, 'true', 'false')", appearance="hidden-answer"); r += 1
    att_intent = (
        "com.example.methodmesh.EXECUTE_METHOD(method_id='attestation.create',caller='odk',"
        "form_id=${mm_form_id},form_version=${mm_form_version},form_instance_id=${mm_workflow_instance_id},"
        "operator_id=${mm_auth_credential_subject_id},input_event_type='odk_form_commitment',"
        "input_event_payload_hash=${mm_event_payload_hash},input_commitment_recipe=${mm_commitment_recipe},"
        "input_attestation_method_version='1.2.0',input_attestation_schema_version='4',"
        "input_verification_method='NfcCredential',input_verification_execution_id=${mm_auth_methodmesh_execution_id},"
        "input_trusted_timestamp=${mm_timestamp_policy},input_payload_mode='FULL',return_mode='flat',"
        "methodmesh_return_namespace='mm_att')"
    )
    add_row(ws, r, "begin group", "mm_create_attestation", "Create MethodMesh attestation",
            hint="This signs the frozen study-data commitment.", appearance="field-list", relevant="string-length(${mm_event_payload_hash}) = 64 and ${mm_current_data_matches_frozen_commitment} = 'true'", intent=att_intent); r += 1
    for field in ATT_FIELDS:
        add_row(ws, r, "text", f"mm_att_{field}", f"MethodMesh attestation return: {field}", appearance="hidden-answer", readonly="yes"); r += 1
    add_row(ws, r, "end group", "mm_create_attestation_end"); r += 1
    add_row(ws, r, "calculate", "mm_ready_to_submit", calculation="if(${mm_att_methodmesh_status} = 'Succeeded' and ${mm_att_event_payload_hash} = ${mm_event_payload_hash} and string-length(${mm_att_attestation_hash}) = 64, 'true', 'false')", appearance="hidden-answer"); r += 1
    add_row(ws, r, "select_one mm_yes", "mm_submission_guard", "Confirm attestation verified", required="yes", relevant="${mm_finalize_for_attestation} = 'yes'", constraint="${mm_ready_to_submit} = 'true' and . = 'yes'", constraint_message="MethodMesh attestation is incomplete or the data no longer match the frozen commitment.")

def add_recipe_sheet(wb):
    if "MethodMesh recipe" in wb.sheetnames:
        del wb["MethodMesh recipe"]
    ws = wb.create_sheet("MethodMesh recipe", 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "MethodMesh Excel-first recipe builder"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:J1")
    ws["A3"] = "How to use"
    ws["A3"].font = Font(bold=True)
    ws["A4"] = "1. Add ordinary study fields to survey rows 30–468."
    ws["A5"] = "2. Enter yes in the Include field? column only for fields that must be included in the attested recipe. Blank means excluded."
    ws["A6"] = "3. Copy Recipe JSON below into survey.mm_commitment_recipe calculation. Copy Canonical expression into survey.mm_canonical_commitment_live calculation."
    ws["A7"] = "4. Replace mm_form_id and mm_form_version calculations with the values from settings, then convert the XLSForm normally."
    ws["A9"] = "Recipe JSON — copy the displayed result from B9 into survey"
    ws["A9"].font = Font(bold=True)
    fixed = '{"schema":"methodmesh.commitment_recipe.v1","canonicalization":"ordered-kv-v1","hash_algorithm":"SHA-256","encoding":"UTF-8","pair_separator":"|","key_value_separator":"=","escaping":"none","members":['
    context = '[{"path":"study_id","type":"string","commitment":"value"},{"path":"form_id","type":"string","commitment":"value"},{"path":"form_version","type":"string","commitment":"value"},{"path":"workflow_instance_id","type":"string","commitment":"value"},{"path":"credential_id_sha256","type":"sha256","commitment":"text-utf8-sha256"},{"path":"credential_subject_id_sha256","type":"sha256","commitment":"text-utf8-sha256"},{"path":"issuer_public_key_fingerprint_sha256","type":"sha256","commitment":"value"},{"path":"verification_evidence_hash","type":"sha256","commitment":"value"}'
    ws["A8"] = "Recipe prefix (generated; do not edit)"
    ws["B8"] = fixed + context[1:]
    ws["B9"] = '=B8&K451&"]}"'
    ws["A10"] = "Canonical expression — copy the displayed result from B10 into survey"
    ws["A10"].font = Font(bold=True)
    ws["B10"] = '="concat(\'study_id=\', string(${mm_study_id}), \'|\', \'form_id=\', string(${mm_form_id}), \'|\', \'form_version=\', string(${mm_form_version}), \'|\', \'workflow_instance_id=\', string(${mm_workflow_instance_id}), \'|\', \'credential_id_sha256=\', string(${mm_credential_id_sha256}), \'|\', \'credential_subject_id_sha256=\', string(${mm_credential_subject_id_sha256}), \'|\', \'issuer_public_key_fingerprint_sha256=\', string(${mm_auth_issuer_public_key_fingerprint_sha256}), \'|\', \'verification_evidence_hash=\', string(${mm_auth_verification_evidence_hash}), \'|\',"&L455&")"'
    ws["B10"].value = ws["B10"].value.replace("L455", "L451")
    b10_prefix, _, b10_suffix = ws["B10"].value.rpartition(", '|',")
    ws["B10"].value = b10_prefix + "," + b10_suffix
    headers = ["survey row", "name", "type", "Include field?", "decision", "recipe type", "transform", "recipe member JSON", "canonical pair", "notes", "recipe accumulation", "canonical accumulation"]
    for col, value in enumerate(headers, 1):
        ws.cell(RECIPE_START - 1, col).value = value
        ws.cell(RECIPE_START - 1, col).font = Font(bold=True, color="FFFFFF")
        ws.cell(RECIPE_START - 1, col).fill = PatternFill("solid", fgColor="5B9BD5")
    for i, source_row in enumerate(range(SURVEY_START, SURVEY_END + 1), RECIPE_START):
        ws.cell(i, 1).value = source_row
        ws.cell(i, 2).value = f'=survey!B{source_row}'
        ws.cell(i, 3).value = f'=survey!A{source_row}'
        ws.cell(i, 4).value = ""
        ws.cell(i, 5).value = f'=IF(OR(B{i}="",C{i}=""),"",IF(LOWER(D{i})="yes","INCLUDE","EXCLUDE"))'
        ws.cell(i, 6).value = f'=IF(E{i}<>"INCLUDE","",IF(LEFT(LOWER(C{i}),10)="select_one ","select_one",IF(LOWER(C{i})="date","date",IF(OR(LOWER(C{i})="integer",LOWER(C{i})="decimal",LOWER(C{i})="range"),LOWER(C{i}),"sha256"))))'
        ws.cell(i, 7).value = f'=IF(E{i}<>"INCLUDE","",IF(F{i}="date","yyyy-mm-dd",IF(F{i}="select_one","stored-choice-name",IF(F{i}="sha256","odk-lexical-utf8-sha256","odk-canonical-scalar"))))'
        ws.cell(i, 8).value = f'=IF(E{i}<>"INCLUDE","","{{"&CHAR(34)&"path"&CHAR(34)&":"&CHAR(34)&B{i}&CHAR(34)&","&CHAR(34)&"type"&CHAR(34)&":"&CHAR(34)&F{i}&CHAR(34)&","&CHAR(34)&"xlsform_type"&CHAR(34)&":"&CHAR(34)&C{i}&CHAR(34)&","&CHAR(34)&"commitment"&CHAR(34)&":"&CHAR(34)&IF(F{i}="sha256","text-utf8-sha256","value")&CHAR(34)&","&CHAR(34)&"transform"&CHAR(34)&":"&CHAR(34)&G{i}&CHAR(34)&"}}")'
        name_ref = f'&B{i}&'
        date_expr = '"if(string-length(string(${\"'+name_ref+'\"})) > 0, format-date(${\"'+name_ref+'\"}, \'%Y-%m-%d\'), \'\')"'
        hash_expr = '"digest(string(${\"'+name_ref+'\"}), \'SHA-256\', \'hex\')"'
        value_expr = '"string(${\"'+name_ref+'\"})"'
        ws.cell(i, 9).value = f'=IF(E{i}<>"INCLUDE","","\'|"&B{i}&"=\', "&IF(F{i}="date",{date_expr},IF(F{i}="sha256",{hash_expr},{value_expr})))'
        ws.cell(i, 10).value = f'=IF(E{i}="INCLUDE","Included by explicit yes","Not included")'
    ws.cell(RECIPE_START, 11).value = '=IF(H13="","",","&H13)'
    ws.cell(RECIPE_START, 12).value = '=IF(I13="","",I13)'
    for i in range(RECIPE_START + 1, RECIPE_START + (SURVEY_END - SURVEY_START + 1)):
        ws.cell(i, 11).value = f'=K{i-1}&IF(H{i}="","",","&H{i})'
        ws.cell(i, 12).value = f'=L{i-1}&IF(I{i}="","",","&I{i})'
    for col, width in {"A":12,"B":28,"C":24,"D":16,"E":14,"F":16,"G":30,"H":90,"I":100,"J":24}.items(): ws.column_dimensions[col].width = width
    ws.column_dimensions["K"].width = 90
    ws.column_dimensions["L"].width = 100
    dv = DataValidation(type="list", formula1='"yes"', allow_blank=True)
    dv.error = "Enter yes to include this field in the recipe, or leave it blank to exclude it."
    dv.errorTitle = "MethodMesh recipe inclusion"
    ws.add_data_validation(dv)
    dv.add(f"D{RECIPE_START}:D{RECIPE_START + (SURVEY_END - SURVEY_START)}")
    ws.conditional_formatting.add(
        f"D{RECIPE_START}:D{RECIPE_START + (SURVEY_END - SURVEY_START)}",
        FormulaRule(formula=[f'$D{RECIPE_START}="yes"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    ws.freeze_panes = "A13"
    ws["B9"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B10"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[9].height = 80; ws.row_dimensions[10].height = 100


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copyfile(SOURCE, OUTPUT)
    wb = load_workbook(OUTPUT)
    survey = wb["survey"]
    add_generic_scaffold(survey)
    if not any(wb["choices"].cell(r, 1).value == "mm_yes" and wb["choices"].cell(r, 2).value == "yes" for r in range(2, wb["choices"].max_row + 1)):
        wb["choices"].append(["mm_yes", "yes", "Yes"])
    add_recipe_sheet(wb)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
