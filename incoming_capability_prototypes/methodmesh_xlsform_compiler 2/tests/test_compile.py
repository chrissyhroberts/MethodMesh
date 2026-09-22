import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from methodmesh_xlsform.compiler import compile_xlsform
from methodmesh_xlsform.errors import ValidationError


class CompileTests(unittest.TestCase):
    def make_source(self, path: Path):
        wb = Workbook()
        ws = wb.active
        ws.title = "survey"
        ws.append(["type", "name", "label", "readonly", "mm_commit"])
        ws.append(["text", "participant_id", "Participant ID", None, None])
        ws.append(["integer", "age", "Age", "no", None])
        ws.append(["select_one outcome", "outcome", "Outcome", None, None])
        ws.append(["text", "ui_helper", "UI helper", None, "exclude"])
        ws.append(["calculate", "calc", None, None, None])
        ws["C6"] = None
        ws["A6"] = "calculate"
        ws["B6"] = "calc"
        # calculation column may be absent in source; compiler should add it.

        ch = wb.create_sheet("choices")
        ch.append(["list_name", "name", "label"])
        ch.append(["outcome", "normal", "Normal"])
        ch.append(["outcome", "abnormal", "Abnormal"])

        st = wb.create_sheet("settings")
        st.append(["form_title", "form_id", "version"])
        st.append(["Test", "test_form", "2026092201"])

        mm = wb.create_sheet("methodmesh")
        mm.append(["key", "value"])
        mm.append(["study_id", "TEST_STUDY"])
        mm.append(["timestamp_policy", "preferred"])
        wb.save(path)

    def test_full_compile(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            result = compile_xlsform(source, out)
            self.assertTrue(result.output_xlsx.exists())
            release = load_workbook(result.output_xlsx, data_only=False)
            self.assertNotIn("methodmesh", release.sheetnames)
            ws = release["survey"]
            headers = [c.value for c in ws[1]]
            self.assertNotIn("mm_commit", headers)
            names = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
            self.assertIn("mm_authenticate_operator", names)
            self.assertIn("mm_create_attestation", names)
            self.assertIn("mm_submission_guard", names)
            participant_row = names.index("participant_id") + 2
            relevant_col = headers.index("relevant") + 1
            readonly_col = headers.index("readonly") + 1
            self.assertIn("mm_auth_ok", ws.cell(participant_row, relevant_col).value)
            self.assertIn("mm_finalize_for_attestation", ws.cell(participant_row, readonly_col).value)
            age_row = names.index("age") + 2
            self.assertNotIn("no", ws.cell(age_row, readonly_col).value.lower())
            manifest = json.loads(result.manifest_json.read_text())
            self.assertEqual([m["name"] for m in manifest["commitment"]["members"]], ["participant_id", "age", "outcome"])
            explicit = [x["name"] for x in manifest["exclusions"] if x["explicit"]]
            self.assertEqual(explicit, ["ui_helper"])

    def test_repeat_requires_explicit_exclusion(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            self.make_source(source)
            wb = load_workbook(source)
            ws = wb["survey"]
            ws.insert_rows(2, 3)
            ws.cell(2, 1, "begin_repeat"); ws.cell(2, 2, "people")
            ws.cell(3, 1, "text"); ws.cell(3, 2, "person_name"); ws.cell(3, 3, "Name")
            ws.cell(4, 1, "end_repeat")
            wb.save(source)
            with self.assertRaises(ValidationError):
                compile_xlsform(source, td / "out")


if __name__ == "__main__":
    unittest.main()
